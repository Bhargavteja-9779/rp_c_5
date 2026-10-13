"""Minimal block-based DOCX writer for an Elsevier (Analytica Chimica Acta) submission.

Blocks: ("title", s) ("h1", s) ("h2", s) ("p", s) ("plabel", label, s) ("eq", latex, number)
("fig", png_path, caption) ("table", DataFrame, caption, note) ("bullets", [s]) ("pagebreak",)
("refs", [s]). Inline markup in paragraphs: *italic*, **bold**, ^{sup}, _{sub}.
"""
from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

FONT = "Times New Roman"
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\})")


def render_equation(latex: str, path: Path, fontsize: int = 13) -> Path:
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${latex}$", fontsize=fontsize, math_fontfamily="stix")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.03, transparent=False)
    plt.close(fig)
    return path


def _runs(par, text: str, size: float | None = None, bold: bool = False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True; r.bold = bold
        elif part.startswith("^{"):
            r = par.add_run(part[2:-1]); r.font.superscript = True; r.bold = bold
        elif part.startswith("_{"):
            r = par.add_run(part[2:-1]); r.font.subscript = True; r.bold = bold
        else:
            r = par.add_run(part); r.bold = bold
        r.font.name = FONT
        if size:
            r.font.size = Pt(size)


def _line_numbers(section):
    sect = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
    sect.append(ln)


def _page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)


def _cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def build_docx(blocks: list, out: Path, tmp: Path, line_numbers: bool = True) -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    if line_numbers:
        _line_numbers(sec)
    _page_numbers(sec)
    eq_dir = tmp / "eq"; eq_dir.mkdir(parents=True, exist_ok=True)
    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, b[1], size=16, bold=True)
        elif kind == "center":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _runs(p, b[1])
        elif kind in ("h1", "h2", "h3"):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12 if kind == "h1" else 8)
            p.paragraph_format.keep_with_next = True
            _runs(p, b[1], size=13 if kind == "h1" else 12, bold=kind != "h3")
            if kind == "h3":
                for r in p.runs:
                    r.italic = True
        elif kind == "p":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; _runs(p, b[1])
        elif kind == "plabel":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(b[1] + " "); r.bold = True; r.font.name = FONT
            _runs(p, b[2])
        elif kind == "bullets":
            for s in b[1]:
                p = doc.add_paragraph(style="List Bullet"); _runs(p, s)
        elif kind == "numbered":
            for i, s in enumerate(b[1], 1):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
                _runs(p, f"({i})\t{s}")
        elif kind == "eq":
            img = render_equation(b[1], eq_dir / f"eq{b[2]}.png")
            p = doc.add_paragraph(); p.paragraph_format.keep_together = True
            ts = p.paragraph_format.tab_stops
            ts.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t")
            from PIL import Image

            w, h = Image.open(img).size
            width_cm = min(13.0, w / 300 * 2.54)
            p.add_run().add_picture(str(img), width=Cm(width_cm))
            r = p.add_run(f"\t({b[2]})"); r.font.name = FONT
        elif kind == "fig":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(b[1]), width=Cm(b[3] if len(b) > 3 else 16.0))
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            c.paragraph_format.line_spacing = 1.15
            _runs(c, b[2], size=10)
        elif kind == "table":
            df, caption, note = b[1], b[2], b[3] if len(b) > 3 else ""
            c = doc.add_paragraph(); c.paragraph_format.keep_with_next = True; c.paragraph_format.line_spacing = 1.15
            _runs(c, caption, size=10)
            t = doc.add_table(rows=1, cols=len(df.columns))
            t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, col in enumerate(df.columns):
                cell = t.rows[0].cells[j]; cell.text = ""
                _runs(cell.paragraphs[0], str(col), size=8, bold=True); _cell_shading(cell, "E8E8E8")
            for _, row in df.iterrows():
                cells = t.add_row().cells
                for j, v in enumerate(row.values):
                    cells[j].text = ""
                    _runs(cells[j].paragraphs[0], str(v), size=8)
            for row in t.rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(0)
            if note:
                n = doc.add_paragraph(); n.paragraph_format.line_spacing = 1.15; _runs(n, note, size=9)
        elif kind == "refs":
            for i, s in enumerate(b[1], 1):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
                p.paragraph_format.line_spacing = 1.15
                _runs(p, f"[{i}]\t{s}", size=10)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(out)
